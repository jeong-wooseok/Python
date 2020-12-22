from openpyxl import load_workbook
wb = load_workbook('sample.xlsx')
ws = wb.active
ws.title = '수정본'

from openpyxl.chart import BarChart, Reference , LineChart

# 1. Barchart
bar_value = Reference(ws, min_row=1, max_row=11, min_col=2,max_col=3)
bar_chart = BarChart() # 차트 종류 설정
bar_chart.add_data(bar_value,titles_from_data=True)
ws.add_chart(bar_chart, 'E1')

# 2. LineChart
line_value = Reference(ws, min_row=1, max_row=11, min_col=2,max_col=3)
line_chart = LineChart() # 차트 종류 설정
line_chart.add_data(line_value,titles_from_data=True)
line_chart.title = '성적표'
line_chart.style = '20'
line_chart.y_axis ='점수'
line_chart.x_axis = '번호'
ws.add_chart(line_chart, 'E15')

wb.save("sample_modi.xlsx")